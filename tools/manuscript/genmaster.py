# サイト（index.html）を正として原稿マスター docx を生成する。
# 使い方: node dump.mjs && python3 genmaster.py [出力先.docx]
import json, datetime, re, os, sys
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = os.path.dirname(os.path.abspath(__file__))
SRC  = os.path.join(HERE, 'site.json')
OUT  = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, '原稿マスター_ポートフォリオサイト.docx')
D = json.load(open(SRC, encoding='utf-8'))
doc = Document()

st = doc.styles['Normal']
st.font.name = 'Yu Gothic'
st.font.size = Pt(10.5)

def P(t='', style=None, size=None, bold=False, color=None, space_after=6, indent=0):
    p = doc.add_paragraph(style=style)
    r = p.add_run(t)
    if size: r.font.size = Pt(size)
    if bold: r.bold = True
    if color: r.font.color.rgb = RGBColor(*color)
    p.paragraph_format.space_after = Pt(space_after)
    if indent: p.paragraph_format.left_indent = Pt(indent)
    return p

def H(t, lv): doc.add_heading(t, level=lv)
def label(t): P(t, size=9, bold=True, color=(0x88,0x88,0x88), space_after=2)

def clean(s): return re.sub(r'\s+', ' ', (s or '')).strip()

# ---------- 表紙 ----------
H('原稿マスター｜ポートフォリオサイト', 0)
P(f"生成日：{datetime.date.today():%Y年%m月%d日}", size=9, color=(0x88,0x88,0x88), space_after=2)
P('出典：index.html（サイトが正。この原稿はサイトから生成したもの）',
  size=9, color=(0x88,0x88,0x88), space_after=18)

# ---------- 1. 共通 ----------
H('1. 共通部分', 1)
H('1-1. ヒーロー', 2)
if D['hero']['lead']: P(clean(D['hero']['lead']))
H('1-2. コンセプト図（ヒーローの図版内テキスト）', 2)
P('　／　'.join(clean(t) for t in D['concept']['svg']), space_after=12)

H('1-3. プロジェクトカード（6枚）', 2)
for i, c in enumerate(D['hero']['cards'], 1):
    if c.get('cat'): label(clean(c['cat']))
    P(clean(c['t']), bold=True, space_after=2)
    P(clean(c['lead']), space_after=10)

H('1-4. About（経歴と実績）', 2)
P(clean(D['about']['head']), bold=True, space_after=2)
P(f"{clean(D['about']['name'])}　／　{clean(D['about']['kana'])}", space_after=8)
label('年表')
for e in D['about']['timeline']:
    P(f"{clean(e['y'])}　{clean(e['t'])}", space_after=2, indent=12)
if D['about']['now']: P(clean(D['about']['now']), space_after=8, indent=12)
for t in D['about']['body']: P(clean(t))

H('1-5. フッター', 2)
P(clean(D['foot']['lead']))
if D['foot']['mail']: P('連絡先：' + D['foot']['mail'].replace('mailto:', ''), size=9)
P(clean(D['foot']['copy']), size=9, color=(0x88,0x88,0x88))

# ---------- 2. ケース詳細 ----------
H('2. ケース詳細', 1)

def cells_line(cells):
    if not cells: return None
    cs = [('［動画］' + c[3:]) if c.startswith('動画：') else c for c in cells]
    return 'スライダー（{}枚）：{}'.format(len(cs), ' ／ '.join(cs))

def imgs_line(imgs):
    if not imgs: return None
    out = []
    for x in imgs:
        out.append(x if isinstance(x, str) else clean(x.get('cap') or x.get('f')))
    return '画像（{}枚）：{}'.format(len(out), ' ／ '.join(out))

def write_case(key):
    m = D['meta'][key]
    H(f"{m['no']} — {clean(m['t'])}", 2)
    c = D['c01'] if key == 'c01' else D['cases'][key]

    H('Context / 課題', 3)
    ctx = c['context']
    if ctx.get('cap'): label('［課題画像］' + clean(ctx['cap']))
    elif not ctx.get('f') and not ctx.get('img'): label('［課題画像なし（意図的）］')
    for t in ctx['paras']: P(clean(t))

    H('Judgment / 判断', 3)
    for i, j in enumerate(c['judgment'], 1):
        P(f"Point {i:02d}｜{clean(j['h'])}", bold=True, space_after=2)
        P(clean(j['t']), space_after=10)

    H('Evidence / 実物', 3)
    steps = c.get('steps') or []
    for si, e in enumerate(c['evidence']):
        if si < len(steps):
            st_ = steps[si]
            P(f"{clean(st_['cat'])}｜{clean(st_['h'])}", bold=True, space_after=2)
            if st_.get('cap'): P(clean(st_['cap']), size=9, color=(0x55,0x55,0x55), space_after=4)
        head = clean(e.get('h'))
        no = clean(e.get('no'))
        if not steps:
            P((f'（{no}）' if no else '') + head, bold=True, space_after=2)
        meta = ' ／ '.join(x for x in [clean(e.get('sub')), clean(e.get('stack'))] if x)
        if meta: P(meta, size=9, color=(0x55,0x55,0x55), space_after=2)
        line = cells_line(e.get('cells')) or imgs_line(e.get('imgs'))
        if line: label(line)
        for t in (e['paras'] if 'paras' in e else [e.get('t', '')]):
            if clean(t): P(clean(t), space_after=4)
        if e.get('link'): P('外部リンク：' + e['link'], size=9, space_after=10)
        else: doc.paragraphs[-1].paragraph_format.space_after = Pt(10)

    H('Result / 成果', 3)
    P(clean(c['result']['head']), bold=True, space_after=4)
    for t in c['result']['paras']: P(clean(t))

for key in D['order']:
    write_case(key)

# ---------- 3. All Works ----------
H('3. All Works（{}件）'.format(len(D['works'])), 1)
P(clean(D['awHead']['h']), bold=True, space_after=2)
P(clean(D['awHead']['lead']), space_after=12)
for key in D['order']:
    m = D['meta'][key]
    ws = [w for w in D['works'] if w['c'] == key]
    if not ws: continue
    H(f"{m['no']} — {clean(m['t'])}", 2)
    for w in ws:
        P(clean(w['t']), bold=True, space_after=2)
        P('　／　'.join([w.get('y',''), w.get('cat',''), '担当：' + w.get('role','')]),
          size=9, color=(0x55,0x55,0x55), space_after=2)
        P(clean(w.get('desc')), space_after=2)
        il = imgs_line(w.get('imgs'))
        if il: label(il)
        if w.get('video'): label('動画：' + w['video'])
        if w.get('yt'): label('YouTube：' + w['yt'])
        if w.get('link'): P('外部リンク：' + w['link'], size=9, space_after=10)
        else: doc.paragraphs[-1].paragraph_format.space_after = Pt(10)

# ---------- 4. 取りこぼし検出 ----------
written = set()
for p in doc.paragraphs:
    written.add(re.sub(r'[\s　]', '', p.text))
def covered(t):
    n = re.sub(r'[\s　]', '', t)
    if not n: return True
    return any(n in w for w in written)
rest = [t for t in D['leaves'] if not covered(t) and len(t) > 3]
if rest:
    H('4. 本文に拾われていない固定文言', 1)
    P('ナビゲーション、ボタン、ラベルなど。原稿としては管理対象外だが、差分確認のために列挙する。',
      size=9, color=(0x88,0x88,0x88), space_after=8)
    for t in rest: P('・' + clean(t), size=9, space_after=2)

doc.save(OUT)
print('保存', OUT, '／ 段落数', len(doc.paragraphs), '／ 本文に拾われていない固定文言', len(rest))
